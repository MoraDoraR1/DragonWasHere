"""기준 문서 표와 재화배치 엑셀에서 Docs/Data/*.csv를 생성한다.

생성 대상 (나머지 CSV는 Docs/Data/README.md의 출처를 보고 손으로 관리한다)
- stages.csv            : Docs/임시unity제작.md 3.2 표
- air_units.csv         : Docs/임시unity제작.md 4.5 「스테이지별 공중 유닛」 표
- drop_composition.csv  : 엑셀 「드랍 구성 (개편)」 시트
- air_drop.csv          : 엑셀 「공중 유닛 드랍」 시트

사용: python Tools/export_data.py
CSV 머리글은 Unity 데이터 클래스의 필드 이름(snake_case)과 같다. CsvImporter가 이름으로 연결한다.
"""
import csv
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
GUIDE = ROOT / "Docs" / "임시unity제작.md"
EXCEL = ROOT / "Docs" / "기획서" / "20261007_레벨디자인_재화배치_임시값_Ver06.xlsx"
OUT = ROOT / "Docs" / "Data"

HURDLES = {"튜토리얼": "Tutorial", "돌파": "Breakthrough", "소허들": "MinorHurdle", "평탄": "Flat", "대허들": "MajorHurdle"}
BUILDING_TYPES = {"일반": "Normal", "공격": "Attack", "랜드마크": "Landmark", "황금 (보너스)": "Golden"}
AIR_SLOTS = {"1번 유닛": 1, "2번 유닛": 2}
# 엑셀 드랍 구성 시트의 재화 열 순서 (왕관 → 금 코인)
RESOURCE_COLUMNS = ["crown", "diamond", "ruby", "sapphire", "gold_bar", "gold_coin"]


def md_table_after(heading: str) -> list[list[str]]:
    """heading 줄 다음에 처음 나오는 마크다운 표의 데이터 행을 돌려준다."""
    lines = GUIDE.read_text(encoding="utf-8").splitlines()
    start = next(i for i, l in enumerate(lines) if l.startswith(heading))
    rows, in_table = [], False
    for line in lines[start + 1:]:
        if line.startswith("|"):
            in_table = True
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if not all(re.fullmatch(r"-+", c) for c in cells):
                rows.append(cells)
        elif in_table:
            break
    return rows[1:]  # 머리글 제외


def num(text: str):
    if text in ("-", "—", ""):
        return 0
    return float(text) if "." in text else int(text)


def write(name: str, header: list[str], rows: list[list]):
    path = OUT / name
    with path.open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(header)
        w.writerows(rows)
    print(f"{name}: {len(rows)}행")


def export_stages():
    rows = []
    for index, c in enumerate(md_table_after("### 3.2 스테이지별 수치"), 1):
        era_name, theme_name = c[1].split("/")
        rows.append([
            c[0], index, (index - 1) // 4 + 1, (index - 1) // 2 + 1, era_name, theme_name,
            num(c[2]), num(c[3]), num(c[4]),
            num(c[5]), num(c[6]), num(c[7]),
            num(c[8]), num(c[9]),
            num(c[10]), num(c[11]),
            HURDLES[c[12]], num(c[13]), num(c[14]), num(c[15]),
        ])
    write("stages.csv", [
        "stage_id", "index", "era", "theme", "era_name", "theme_name",
        "normal_count", "attack_count", "landmark_count",
        "normal_hp", "attack_hp", "landmark_hp",
        "attack_damage", "landmark_damage",
        "golden_count", "golden_hp",
        "hurdle", "target_repeat", "upgrade_gold_guide", "expected_gold",
    ], rows)


def export_air_units():
    rows = []
    for c in md_table_after("**스테이지별 공중 유닛"):
        stage = c[0]
        rows.append([stage, 1, c[1], num(c[2]), num(c[3]), num(c[4]), num(c[5])])
        if c[6] != "-":
            rows.append([stage, 2, c[6], num(c[7]), num(c[8]), num(c[9]), num(c[10])])
    write("air_units.csv", ["stage_id", "slot", "name", "count", "hp", "damage", "gold"], rows)


def excel_drop_rows(sheet, key_map: dict, key_name: str):
    ws = openpyxl.load_workbook(EXCEL, data_only=True)[sheet]
    header = [c.value for c in ws[4]]
    col = {name: header.index(name) for name in header if name}
    rows = []
    for r in ws.iter_rows(min_row=5, values_only=True):
        if r[0] not in key_map:
            continue
        counts = [r[col[k]] for k in ["왕관", "다이아몬드", "루비", "사파이어", "금괴", "금 코인"]]
        rows.append([key_map[r[0]], r[col["스테이지"]], r[col["목표 골드"]], r[col["목표 개수"]], *counts, r[col["확정 골드"]]])
    return [key_name, "stage_id", "target_gold", "target_count", *RESOURCE_COLUMNS, "total_gold"], rows


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    export_stages()
    export_air_units()
    header, rows = excel_drop_rows("드랍 구성 (개편)", BUILDING_TYPES, "building_type")
    write("drop_composition.csv", header, rows)
    header, rows = excel_drop_rows("공중 유닛 드랍", AIR_SLOTS, "slot")
    write("air_drop.csv", header, rows)


if __name__ == "__main__":
    sys.exit(main())
